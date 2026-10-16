#!/usr/bin/env python3
"""Builds every derived CSV the Ruby loaders need, from the institute's raw
source files alone - so today's whole database state can be regenerated
from source files + scripts only, never from any session-specific scratch
file. Run this before personnel_loader.rb / activos_consolidation_loader.rb.

Usage: python3 build_derived_inputs.py <source_dir> <output_dir>

Reads from <source_dir>:
  personal_CBGP_FINAL_MERGED.csv       (ISO-8859-1, semicolon-delimited)
  usuarios_historico_cbgp_*.xlsx       (first matching file)
  usuarios_activos_cbgp_*.xlsx         (first matching file)
  CBGP. ORCID Codes.xlsx

Writes to <output_dir>:
  activos_export.csv              - usuarios_activos, cleaned to plain CSV
  category_dictionary.csv         - Spanish "Categoría Según Reglamento"
                                     label -> ontology member_category code,
                                     derived empirically by joining
                                     usuarios_historico to the already-
                                     ontology-coded personal_CBGP_FINAL_
                                     MERGED on (DNI, start date), majority
                                     vote per label (see project memory
                                     project_activos_historico_consolidation
                                     for why this beats guessing a
                                     translation by hand)
  research_area_dictionary.csv    - same technique, "Grupo" label ->
                                     member_research_area_group code
  orcid_codes.csv                 - CBGP ORCID Codes.xlsx, cleaned to plain
                                     CSV (surname, given, orcid)
"""
import csv
import datetime
import glob
import re
import sys
from collections import Counter, defaultdict

import openpyxl


def norm_dni(value):
    return re.sub(r'[\s-]', '', str(value or '')).upper()


def clean(value):
    if value is None:
        return ''
    return str(value).replace('_x000D_', '\n').strip()


def first_match(source_dir, pattern):
    matches = sorted(glob.glob(f'{source_dir}/{pattern}'))
    if not matches:
        sys.exit(f'no file matching {pattern!r} found in {source_dir}')
    return matches[0]


def parse_export_date(value):
    value = (value or '').strip()
    if not value:
        return None
    try:
        return datetime.datetime.strptime(value, '%d/%m/%Y').date()
    except ValueError:
        return None


def build_activos_export(source_dir, output_dir):
    path = first_match(source_dir, 'usuarios_activos_cbgp_*.xlsx')
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    header = [str(h).strip() if h else '' for h in rows[0]]

    out_path = f'{output_dir}/activos_export.csv'
    with open(out_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for row in rows[1:]:
            writer.writerow([
                v.strftime('%Y-%m-%d') if isinstance(v, datetime.datetime) else clean(v)
                for v in row
            ])
    print(f'wrote {len(rows) - 1} rows to {out_path}')


def build_label_dictionaries(source_dir, output_dir):
    historico_path = first_match(source_dir, 'usuarios_historico_cbgp_*.xlsx')
    wb = openpyxl.load_workbook(historico_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    hist_rows = list(ws.iter_rows(min_row=2, values_only=True))
    hist_idx = defaultdict(list)
    for r in hist_rows:
        dni = norm_dni(r[5])
        start = r[9].date() if hasattr(r[9], 'date') else None
        # columns: 11=Categoria, 12=Categoria Segun Reglamento, 14=Grupo
        hist_idx[(dni, start)].append((r[12], r[14]))

    merged_path = f'{source_dir}/personal_CBGP_FINAL_MERGED.csv'
    with open(merged_path, encoding='iso-8859-1', newline='') as f:
        merged_rows = list(csv.reader(f, delimiter=';', quotechar='"'))
    header = merged_rows[0]
    dni_i = header.index('member_dni_nie_pas')
    start_i = header.index('member_start_date')
    cat_i = header.index('member_category')
    rag_i = header.index('member_research_area_group')

    cat_dict = defaultdict(Counter)
    rag_dict = defaultdict(Counter)
    matched = 0
    for r in merged_rows[1:]:
        dni = norm_dni(r[dni_i])
        start = parse_export_date(r[start_i])
        key = (dni, start)
        if key not in hist_idx:
            continue
        categoria_reg, grupo = hist_idx[key][0]
        if r[cat_i] and categoria_reg:
            cat_dict[clean(categoria_reg)][r[cat_i]] += 1
        if r[rag_i] and grupo:
            rag_dict[clean(grupo)][r[rag_i]] += 1
        matched += 1

    cat_out = f'{output_dir}/category_dictionary.csv'
    with open(cat_out, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['label', 'code'])
        for label, counts in cat_dict.items():
            writer.writerow([label, counts.most_common(1)[0][0]])

    rag_out = f'{output_dir}/research_area_dictionary.csv'
    with open(rag_out, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['label', 'code'])
        for label, counts in rag_dict.items():
            writer.writerow([label, counts.most_common(1)[0][0]])

    print(f'matched {matched} rows; wrote {len(cat_dict)} category labels to {cat_out}, '
          f'{len(rag_dict)} research-area labels to {rag_out}')


def build_orcid_codes(source_dir, output_dir):
    path = f'{source_dir}/CBGP. ORCID Codes.xlsx'
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=3, values_only=True))

    out_path = f'{output_dir}/orcid_codes.csv'
    n = 0
    with open(out_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['surname', 'given', 'orcid'])
        for r in rows:
            if r[0] and r[2]:
                writer.writerow([clean(r[0]), clean(r[1]), clean(r[2])])
                n += 1
    print(f'wrote {n} rows to {out_path}')


def main():
    if len(sys.argv) != 3:
        sys.exit(f'usage: {sys.argv[0]} <source_dir> <output_dir>')
    source_dir, output_dir = sys.argv[1], sys.argv[2]

    build_activos_export(source_dir, output_dir)
    build_label_dictionaries(source_dir, output_dir)
    build_orcid_codes(source_dir, output_dir)


if __name__ == '__main__':
    main()
